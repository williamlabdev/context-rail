"""Generate the planning DOCX from its canonical Markdown.

Run with the Python executable returned by load_workspace_dependencies.
Render the result with the documents skill renderer before distribution.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


LATIN_FONT = "Arial"
# Use a font with Traditional Chinese coverage in the bundled headless
# LibreOffice runtime.
CJK_FONT = "Arial Unicode MS"
CONTENT_WIDTH = 7.0


def set_font(run, size=None, bold=None, color="000000"):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), CJK_FONT)
    # Tell LibreOffice that CJK runs should use the eastAsia face. Without
    # this hint its headless renderer may apply Arial to Traditional Chinese
    # and emit missing-glyph boxes in the exported QA pages.
    rfonts.set(qn("w:hint"), "eastAsia")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-TW")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def xml_child(parent, tag, attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    parent.append(node)
    return node


def hyperlink(paragraph, label, url):
    node = OxmlElement("w:hyperlink")
    if url.startswith(("https://", "http://")):
        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        node.set(qn("r:id"), rid)
    else:
        # Keep local references readable in an exported document without
        # tying the file to a particular workstation's path.
        paragraph.add_run(label)
        return
    run = OxmlElement("w:r")
    props = xml_child(run, "w:rPr", {})
    xml_child(props, "w:rFonts", {"w:ascii": LATIN_FONT, "w:hAnsi": LATIN_FONT,
                                 "w:eastAsia": CJK_FONT, "w:hint": "eastAsia"})
    xml_child(props, "w:lang", {"w:eastAsia": "zh-TW"})
    xml_child(props, "w:color", {"w:val": "1F4E79"})
    text = xml_child(run, "w:t", {})
    text.text = label
    node.append(run)
    paragraph._p.append(node)


def inline(paragraph, text, size=None, bold=False, color="000000"):
    pattern = r"(\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)"
    for part in re.split(pattern, text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            hyperlink(paragraph, *link.groups())
        else:
            strong = part.startswith("**") and part.endswith("**")
            code = part.startswith("`") and part.endswith("`")
            value = part[2:-2] if strong else part[1:-1] if code else part
            run = paragraph.add_run(value)
            set_font(run, size=size, bold=bold or strong, color=color)


def display_width(text):
    return sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in text)


def wrap_code(text, limit=88):
    result, buf, size = [], "", 0
    for char in text:
        width = display_width(char)
        if size + width > limit:
            result.append(buf)
            buf, size = "    ", 4
        buf += char
        size += width
    result.append(buf)
    return result


def table_widths(rows):
    cols = len(rows[0])
    if cols == 2:
        return [1.65, CONTENT_WIDTH - 1.65]
    if rows[0][0] in ("ID",):
        return [0.65, 1.65, CONTENT_WIDTH - 2.3]
    if rows[0][-1] == "合計":
        return [4.0, 1.0, 1.0, 1.0]
    if rows[0][0] == "預算項目":
        return [2.4, 1.1, 3.5]
    if rows[0][0] == "情境":
        return [0.85, 1.7, 1.7, 2.75]
    if rows[0][0] == "週次":
        return [0.75, 1.15, 2.65, 2.45]
    if cols == 3:
        return [1.45, 2.5, 3.05]
    return [1.2] + [(CONTENT_WIDTH - 1.2) / (cols - 1)] * (cols - 1)


def make_table(doc, rows):
    widths = table_widths(rows)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    props = table._tbl.tblPr
    borders = xml_child(props, "w:tblBorders", {})
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        xml_child(borders, "w:" + side, {"w:val": "single", "w:sz": 4,
                                         "w:color": "D9D9D9"})
    margins = xml_child(props, "w:tblCellMar", {})
    for side in ("top", "bottom", "left", "right"):
        xml_child(margins, "w:" + side, {"w:w": 90, "w:type": "dxa"})
    for index, values in enumerate(rows):
        row = table.add_row()
        tr_props = row._tr.get_or_add_trPr()
        xml_child(tr_props, "w:cantSplit", {})
        if index == 0:
            xml_child(tr_props, "w:tblHeader", {})
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            fill = "23374D" if index == 0 else "F3F5F7" if index % 2 == 0 else "FFFFFF"
            xml_child(cell._tc.get_or_add_tcPr(), "w:shd", {"w:fill": fill})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.paragraph_format.keep_with_next = index == 0
            if col_index == 0 or re.fullmatch(r"[0-9.,% ]+", value):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(p, value, size=10.5, bold=index == 0,
                   color="FFFFFF" if index == 0 else "000000")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.3
    spacer.add_run().font.size = Pt(3)


def configure(doc):
    # The bundled default template includes decorative borders and theme
    # font references. Remove them so direct font choices are authoritative.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for fonts in doc.styles.element.xpath(".//w:rFonts"):
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + key), None)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.footer_distance = Inches(0.25)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3",
                 "List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        rfonts = style.element.get_or_add_rPr().rFonts
        rfonts.set(qn("w:ascii"), LATIN_FONT)
        rfonts.set(qn("w:hAnsi"), LATIN_FONT)
        rfonts.set(qn("w:eastAsia"), CJK_FONT)
        rfonts.set(qn("w:hint"), "eastAsia")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.widow_control = True
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)
    doc.styles["Title"].font.size = Pt(24)
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(13)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        doc.styles[name].paragraph_format.space_after = Pt(4)
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(p.add_run("ContextRail v4.2   |   "), size=8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    p._p.append(field)


def build(source, output):
    content = source.read_text(encoding="utf-8")
    lines = content.splitlines()
    doc = Document()
    configure(doc)
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.extend(wrap_code(lines[i]))
                i += 1
            for j, value in enumerate(code_lines):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0 if j < len(code_lines)-1 else 7)
                p.paragraph_format.line_spacing = 1.08
                p.paragraph_format.keep_with_next = j < len(code_lines)-1
                p.paragraph_format.left_indent = Inches(0.1)
                run = p.add_run(value)
                set_font(run, size=9.5)
                run.font.name = "Courier New"
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [s.strip() for s in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError("Inconsistent table column count")
            make_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            p = doc.add_paragraph(style="Title" if level == 1 else "Heading " + str(level-1))
            inline(p, heading.group(2), bold=True)
        elif line.startswith("- "):
            inline(doc.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            inline(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            inline(p, line[2:])
        else:
            inline(doc.add_paragraph(), line)
        i += 1
    doc.core_properties.title = "ContextRail 參賽與開發計畫書"
    doc.core_properties.subject = "Environment-aware AI change assurance for internal Cloud Run systems"
    doc.core_properties.author = "ContextRail"
    doc.core_properties.keywords = "five weeks; Go; environment topology; promotion policy; AI change assurance; AINE; agent work order; evidence; Gemini; Cloud Run; release"
    doc.core_properties.comments = "Generated from canonical Markdown; source SHA256 " + hashlib.sha256(content.encode()).hexdigest()
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Generated {output}")
    print(f"Source SHA256 {hashlib.sha256(content.encode()).hexdigest()}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
