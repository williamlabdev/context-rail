"""Build the Hero Demo Engineering Decision Pack DOCX from Markdown.

This is a presentation artifact for the example pack. It deliberately keeps
fixture labels visible so the document cannot be mistaken for deployment proof.
"""

from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LATIN_FONT = "Arial"
# Use a font with Traditional Chinese coverage in the bundled headless
# LibreOffice runtime.
CJK_FONT = "Arial Unicode MS"
CONTENT_WIDTH = 7.25


def set_font(run, size=None, bold=None, italic=None, color="000000", font=LATIN_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), font)
    rpr.rFonts.set(qn("w:hAnsi"), font)
    rpr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def xml_child(parent, tag, attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(qn(key), str(value))
    parent.append(node)
    return node


def inline(paragraph, text, size=9.5, color="000000"):
    pattern = r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
    for part in re.split(pattern, text):
        if not part:
            continue
        strong = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        value = (
            part[2:-2]
            if strong
            else part[1:-1]
            if code
            else link.group(1)
            if link
            else part
        )
        run = paragraph.add_run(value)
        set_font(
            run,
            size=size if not code else max(size - 0.3, 7.5),
            bold=strong,
            color="1F4E79" if link else ("5B2C6F" if code else color),
            font="Courier New" if code else LATIN_FONT,
        )


def table_widths(rows):
    cols = len(rows[0])
    if cols == 2:
        return [1.8, CONTENT_WIDTH - 1.8]
    if cols == 3:
        return [1.35, 2.35, CONTENT_WIDTH - 3.7]
    if cols == 4:
        return [1.15, 2.0, 2.0, CONTENT_WIDTH - 5.15]
    if cols == 5:
        return [1.0, 1.35, 1.8, 1.65, CONTENT_WIDTH - 5.8]
    return [CONTENT_WIDTH / cols] * cols


def make_table(doc, rows):
    widths = table_widths(rows)
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)

    props = table._tbl.tblPr
    borders = xml_child(props, "w:tblBorders", {})
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        xml_child(borders, "w:" + side, {"w:val": "single", "w:sz": 4, "w:color": "D9E0E7"})
    margins = xml_child(props, "w:tblCellMar", {})
    for side in ("top", "bottom", "left", "right"):
        xml_child(margins, "w:" + side, {"w:w": 58, "w:type": "dxa"})

    for row_index, values in enumerate(rows):
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        xml_child(tr_pr, "w:cantSplit", {})
        if row_index == 0:
            xml_child(tr_pr, "w:tblHeader", {})
        for col_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if len(rows[0]) == 2 and row_index > 0 and col_index == 0:
                fill = "EAF1F8"
            else:
                fill = "23374D" if row_index == 0 else ("F5F7F9" if row_index % 2 == 0 else "FFFFFF")
            xml_child(cell._tc.get_or_add_tcPr(), "w:shd", {"w:fill": fill})
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            inline(
                paragraph,
                value,
                size=8.1,
                color="FFFFFF" if row_index == 0 else "000000",
            )
            if row_index == 0 or (len(rows[0]) == 2 and col_index == 0):
                for run in paragraph.runs:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(2)
    spacer.add_run().font.size = Pt(2)


def add_code_block(doc, code):
    for line in code.splitlines() or [""]:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 0.95
        paragraph.paragraph_format.keep_together = True
        ppr = paragraph._p.get_or_add_pPr()
        xml_child(ppr, "w:shd", {"w:fill": "F2F4F7"})
        run = paragraph.add_run(line if line else " ")
        set_font(run, size=7.4, color="3B4450", font="Courier New")
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.add_run().font.size = Pt(1)


def configure(doc):
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for fonts in doc.styles.element.xpath(".//w:rFonts"):
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:" + key), None)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.55)
    section.footer_distance = Inches(0.18)

    for name in (
        "Normal",
        "Title",
        "Subtitle",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "List Bullet",
        "List Number",
    ):
        style = doc.styles[name]
        style.font.name = LATIN_FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        rfonts = style.element.get_or_add_rPr().rFonts
        rfonts.set(qn("w:ascii"), LATIN_FONT)
        rfonts.set(qn("w:hAnsi"), LATIN_FONT)
        rfonts.set(qn("w:eastAsia"), CJK_FONT)
        rfonts.set(qn("w:hint"), "eastAsia")
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.widow_control = True

    normal = doc.styles["Normal"]
    normal.font.size = Pt(9.15)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(3)

    title = doc.styles["Title"]
    title.font.size = Pt(21)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles["Subtitle"]
    subtitle.font.size = Pt(8.7)
    subtitle.font.italic = False
    subtitle.paragraph_format.space_after = Pt(5)

    for name, size in (("Heading 1", 12.2), ("Heading 2", 10.4), ("Heading 3", 9.7)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(35, 55, 77)
        style.paragraph_format.space_before = Pt(6 if name == "Heading 1" else 4)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.size = Pt(8.9)
        style.paragraph_format.left_indent = Inches(0.19)
        style.paragraph_format.first_line_indent = Inches(-0.12)
        style.paragraph_format.space_after = Pt(0.5)
        style.paragraph_format.line_spacing = 1.0

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("ContextRail Hero Demo   |   示例資料   |   "), size=7.5, color="6B7280")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def table_rows(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [part.strip() for part in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-+:?", part) for part in parts):
            rows.append(parts)
        index += 1
    if not rows or any(len(row) != len(rows[0]) for row in rows):
        raise ValueError(f"Inconsistent table at line {start + 1}")
    return rows, index


def build(source: Path, output: Path):
    content = source.read_text(encoding="utf-8")
    lines = content.splitlines()
    doc = Document()
    configure(doc)
    index = 0
    in_code = False
    code_lines = []

    while index < len(lines):
        raw = lines[index]
        line = raw.strip()

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(raw)
            index += 1
            continue
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows, index = table_rows(lines, index)
            make_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = len(heading.group(1))
            style = "Title" if level == 1 else "Heading " + str(level - 1)
            paragraph = doc.add_paragraph(style=style)
            inline(paragraph, heading.group(2), size=20 if level == 1 else 11)
        elif line == ">":
            index += 1
            continue
        elif line.startswith("> "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.15)
            paragraph.paragraph_format.right_indent = Inches(0.08)
            paragraph.paragraph_format.space_after = Pt(3)
            ppr = paragraph._p.get_or_add_pPr()
            xml_child(ppr, "w:shd", {"w:fill": "EAF1F8"})
            inline(paragraph, line[2:], size=8.9)
            for run in paragraph.runs:
                run.bold = True
        elif line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            inline(paragraph, line[2:], size=8.9)
        elif re.match(r"^\d+\. ", line):
            # Keep each Markdown list's visible number local. Word's built-in
            # List Number style otherwise continues numbering across sections.
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.19)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            paragraph.paragraph_format.space_after = Pt(0.5)
            inline(paragraph, line, size=8.9)
        else:
            paragraph = doc.add_paragraph()
            inline(paragraph, line, size=9.15)
        index += 1

    if in_code:
        add_code_block(doc, "\n".join(code_lines))

    source_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()
    doc.core_properties.title = "ContextRail Hero Demo 工程決策報告包"
    doc.core_properties.subject = "Example Engineering Decision Pack with role-specific views"
    doc.core_properties.author = "ContextRail"
    doc.core_properties.keywords = "Engineering Decision Record; Agent Context Pack; Cloud Run; release evidence"
    doc.core_properties.comments = "Example-only artifact generated from canonical Markdown; source SHA256 " + source_hash
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Generated {output}")
    print(f"Source SHA256 {source_hash}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())
